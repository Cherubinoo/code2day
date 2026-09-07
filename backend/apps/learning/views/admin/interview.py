"""Views extracted from the original monolithic apps/learning/views.py
(module: admin/interview). Pure code motion — see apps/learning/views/_imports.py
and _shared.py for why this split can't hit a missing-name error."""
from .._imports import *
from .._shared import *


class AdminInterviewTrackListCreateView(APIView):
    """System Admin: list/create Interview Practice tracks — the top-level
    content bank for the Interview Practice module (mirrors Examination
    for Competitive Practice). A department's students see whichever
    track matches Department.interview_track."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        tracks = InterviewTrack.objects.annotate(
            topic_count=Count('topics', distinct=True),
            question_count=Count('topics__questions', distinct=True),
        )
        return Response([
            {"id": t.id, "key": t.key, "name": t.name, "description": t.description,
             "topic_count": t.topic_count, "question_count": t.question_count}
            for t in tracks
        ])

    def post(self, request):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        name = (request.data.get('name') or '').strip()
        if not name:
            return Response({"error": "name is required."}, status=400)
        key = (request.data.get('key') or '').strip().lower() or slugify(name)[:40]
        if not key:
            return Response({"error": "Could not derive a key from this name — provide one explicitly."}, status=400)
        if InterviewTrack.objects.filter(key=key).exists():
            return Response({"error": f"A track with key '{key}' already exists."}, status=400)
        track = InterviewTrack.objects.create(key=key, name=name, description=(request.data.get('description') or '').strip())
        return Response({"id": track.id, "key": track.key, "name": track.name, "description": track.description,
                          "topic_count": 0, "question_count": 0}, status=201)

class AdminInterviewTrackDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def patch(self, request, track_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        track = get_object_or_404(InterviewTrack, id=track_id)
        if 'name' in request.data:
            name = (request.data.get('name') or '').strip()
            if not name:
                return Response({"error": "name cannot be empty."}, status=400)
            track.name = name
        if 'description' in request.data:
            track.description = (request.data.get('description') or '').strip()
        track.save()
        return Response({"id": track.id, "key": track.key, "name": track.name, "description": track.description})

    def delete(self, request, track_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        track = get_object_or_404(InterviewTrack, id=track_id)
        track.delete()
        return Response({"message": "Track deleted"})

class AdminInterviewTrackTreeView(APIView):
    """System Admin: full Track > Topic > Folder(nested) > Question tree
    for one track — the admin management view, mirrors
    _serialize_examination_syllabus's role for Competitive Practice."""
    permission_classes = [IsAuthenticated]

    def get(self, request, track_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        track = get_object_or_404(
            InterviewTrack.objects.prefetch_related(
                'topics__questions',
                'topics__folders__questions',
                'topics__folders__media_items',
                'topics__folders__subfolders__questions',
                'topics__folders__subfolders__media_items',
            ),
            id=track_id,
        )
        return Response(_serialize_interview_track(track))

class AdminInterviewTopicListCreateView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, track_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        track = get_object_or_404(InterviewTrack, id=track_id)
        title = (request.data.get('title') or '').strip()
        if not title:
            return Response({"error": "title is required."}, status=400)
        if track.topics.filter(title__iexact=title).exists():
            return Response({"error": "A topic with this title already exists in this track."}, status=400)
        topic = InterviewTopic.objects.create(track=track, title=title, order=track.topics.count())
        return Response(_serialize_interview_topic(topic), status=201)

class AdminInterviewTopicDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def delete(self, request, topic_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        topic = get_object_or_404(InterviewTopic, id=topic_id)
        topic.delete()
        return Response({"message": "Topic deleted"})

class AdminInterviewFolderListCreateView(APIView):
    """Create a top-level folder directly under a topic."""
    permission_classes = [IsAuthenticated]

    def post(self, request, topic_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        topic = get_object_or_404(InterviewTopic, id=topic_id)
        title = (request.data.get('title') or '').strip()
        if not title:
            return Response({"error": "title is required."}, status=400)
        if topic.folders.filter(parent__isnull=True, title__iexact=title).exists():
            return Response({"error": "A folder with this name already exists in this topic."}, status=400)
        folder = InterviewFolder.objects.create(topic=topic, title=title, order=topic.folders.filter(parent__isnull=True).count())
        return Response(_serialize_interview_folder(folder), status=201)

class AdminInterviewSubfolderListCreateView(APIView):
    """Create a folder nested inside another folder — same idea one level
    down, so Interview Bank folders nest to arbitrary depth like
    Competitive Bank's."""
    permission_classes = [IsAuthenticated]

    def post(self, request, folder_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        parent = get_object_or_404(InterviewFolder, id=folder_id)
        title = (request.data.get('title') or '').strip()
        if not title:
            return Response({"error": "title is required."}, status=400)
        if parent.subfolders.filter(title__iexact=title).exists():
            return Response({"error": "A folder with this name already exists here."}, status=400)
        folder = InterviewFolder.objects.create(topic=parent.topic, parent=parent, title=title, order=parent.subfolders.count())
        return Response(_serialize_interview_folder(folder), status=201)

class AdminInterviewFolderDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def patch(self, request, folder_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        folder = get_object_or_404(InterviewFolder, id=folder_id)
        title = (request.data.get('title') or '').strip()
        if not title:
            return Response({"error": "title cannot be empty."}, status=400)
        siblings = InterviewFolder.objects.filter(topic=folder.topic, parent=folder.parent)
        if siblings.exclude(id=folder.id).filter(title__iexact=title).exists():
            return Response({"error": "A folder with this name already exists here."}, status=400)
        folder.title = title
        folder.save(update_fields=['title'])
        return Response(_serialize_interview_folder(folder))

    def delete(self, request, folder_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        folder = get_object_or_404(InterviewFolder, id=folder_id)
        folder.delete()
        return Response({"message": "Folder deleted"})

class AdminInterviewFolderMediaView(APIView):
    """System Admin: list/upload real media files (images, PDFs, videos)
    attached to one Interview folder — same idea as
    AdminSyllabusFolderMediaView for Competitive Bank's folders."""
    permission_classes = [IsAuthenticated]

    ALLOWED_CONTENT_TYPES = {
        'image/jpeg', 'image/jpg', 'image/png', 'image/gif', 'image/webp', 'image/svg+xml', 'image/heic', 'image/heif',
        'video/mp4', 'video/webm', 'video/quicktime', 'video/x-msvideo', 'video/x-matroska', 'video/3gpp', 'video/x-m4v',
        'application/pdf',
    }
    MAX_UPLOAD_BYTES = 25 * 1024 * 1024        # 25MB — images, PDFs
    MAX_VIDEO_UPLOAD_BYTES = 200 * 1024 * 1024  # 200MB — a several-minute explainer clip

    def get(self, request, folder_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        folder = get_object_or_404(InterviewFolder, id=folder_id)
        return Response([_serialize_interview_folder_media(m) for m in folder.media_items.all()])

    def post(self, request, folder_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        folder = get_object_or_404(InterviewFolder, id=folder_id)

        uploaded = request.FILES.get('file')
        if not uploaded:
            return Response({"error": "No file provided."}, status=400)
        if uploaded.content_type not in self.ALLOWED_CONTENT_TYPES:
            return Response({"error": f"Unsupported file type: {uploaded.content_type}."}, status=400)
        is_video = uploaded.content_type.startswith('video/')
        limit = self.MAX_VIDEO_UPLOAD_BYTES if is_video else self.MAX_UPLOAD_BYTES
        if uploaded.size > limit:
            return Response({"error": f"File too large. Maximum size is {limit // (1024 * 1024)}MB."}, status=400)

        media = InterviewFolderMedia.objects.create(
            folder=folder,
            file=uploaded,
            title=(request.data.get('title') or '').strip(),
            description=(request.data.get('description') or '').strip(),
            content_type=uploaded.content_type,
            order=folder.media_items.count(),
        )
        return Response(_serialize_interview_folder_media(media), status=201)

class AdminInterviewFolderMediaDetailView(APIView):
    """System Admin: rename or delete one uploaded Interview folder media
    file. Delete also removes it from disk, not just the database row."""
    permission_classes = [IsAuthenticated]

    def patch(self, request, media_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        media = get_object_or_404(InterviewFolderMedia, id=media_id)
        update_fields = []
        if 'title' in request.data:
            title = (request.data.get('title') or '').strip()
            if not title:
                return Response({"error": "title cannot be empty."}, status=400)
            media.title = title
            update_fields.append('title')
        if 'description' in request.data:
            media.description = (request.data.get('description') or '').strip()
            update_fields.append('description')
        if update_fields:
            media.save(update_fields=update_fields)
        return Response(_serialize_interview_folder_media(media))

    def delete(self, request, media_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        media = get_object_or_404(InterviewFolderMedia, id=media_id)
        media.file.delete(save=False)
        media.delete()
        return Response({"message": "Media deleted"})

class AdminInterviewQuestionListCreateView(APIView):
    """System Admin: list/create questions for one topic. Optional
    ?folder_id= scopes both GET and POST to one of the topic's folders
    instead of its top-level (unfoldered) question list — same pattern as
    AdminSubtopicQuestionListCreateView."""
    permission_classes = [IsAuthenticated]

    def _resolve_folder(self, topic, folder_id):
        if not folder_id:
            return None, None
        folder = InterviewFolder.objects.filter(id=folder_id, topic=topic).first()
        if not folder:
            return None, Response({"error": "Folder not found in this topic."}, status=404)
        return folder, None

    def get(self, request, topic_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        topic = get_object_or_404(InterviewTopic, id=topic_id)
        folder, error = self._resolve_folder(topic, request.query_params.get('folder_id'))
        if error:
            return error
        questions = folder.questions.all() if folder else topic.questions.filter(folder__isnull=True)
        return Response([_serialize_interview_question(q) for q in questions])

    def post(self, request, topic_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        topic = get_object_or_404(InterviewTopic, id=topic_id)
        folder, error = self._resolve_folder(topic, request.data.get('folder_id'))
        if error:
            return error
        cleaned, error = _validate_interview_question_payload(request.data)
        if error:
            return error
        sibling_count = folder.questions.count() if folder else topic.questions.filter(folder__isnull=True).count()
        q = InterviewQuestion.objects.create(topic=topic, folder=folder, order=sibling_count, **{
            f: cleaned.get(f, '') for f in INTERVIEW_QUESTION_FIELDS if f not in ('question_type', 'difficulty')
        }, question_type=cleaned.get('question_type') or 'conceptual', difficulty=cleaned.get('difficulty') or 'Beginner')
        return Response(_serialize_interview_question(q), status=201)

class AdminInterviewQuestionDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def patch(self, request, question_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        q = get_object_or_404(InterviewQuestion, id=question_id)
        cleaned, error = _validate_interview_question_payload(request.data, require_core=False)
        if error:
            return error
        for field, value in cleaned.items():
            setattr(q, field, value)
        q.save()
        return Response(_serialize_interview_question(q))

    def delete(self, request, question_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        q = get_object_or_404(InterviewQuestion, id=question_id)
        q.delete()
        return Response({"message": "Question deleted"})

class AdminInterviewQuestionBulkUploadView(APIView):
    """System Admin: bulk-populate an Interview Track's Topic > Folder >
    Question tree from an uploaded .xlsx/.xls/.csv, column-name-driven —
    same approach as AdminExaminationSyllabusUploadView. Expects the
    columns from the real upload template: Question ID, Field, Topic,
    Subtopic, Question Type, Difficulty, Question, Answer / Model Answer,
    Follow-up Question, Follow-up Answer, Tools / Technologies, Key
    Concepts / Keywords, Source Reference. "Field" names the track (e.g.
    "Cybersecurity") — resolved/created by name, independent of whatever
    track_id the request was made under, so re-uploading the same sheet
    always lands in the same track regardless of which track tile you
    clicked. "Subtopic" is optional and maps to a top-level Folder."""
    permission_classes = [IsAuthenticated]

    QUESTION_TYPE_LABEL_TO_KEY = {label.lower(): key for key, label in InterviewQuestion.QUESTION_TYPE_CHOICES}
    DIFFICULTY_LABEL_TO_KEY = {label.lower(): key for key, label in InterviewQuestion.DIFFICULTY_CHOICES}

    def post(self, request, track_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        get_object_or_404(InterviewTrack, id=track_id)  # just validates the URL's track exists

        upload = request.FILES.get("file")
        if not upload:
            return Response({"error": "No file uploaded."}, status=400)

        filename = upload.name.lower()
        try:
            if filename.endswith((".xlsx", ".xls")):
                rows = self._read_excel(upload)
            elif filename.endswith(".csv"):
                rows = self._read_csv(upload)
            else:
                return Response({"error": "Only .xlsx, .xls, or .csv files are supported."}, status=400)
        except Exception as e:
            return Response({"error": f"Could not read file: {e}"}, status=400)

        if not rows:
            return Response({"error": "File is empty."}, status=400)

        header = [str(h or '').strip().lower() for h in rows[0]]
        col_idx = {h: i for i, h in enumerate(header) if h}

        def col(row, name):
            idx = col_idx.get(name)
            if idx is None or idx >= len(row) or row[idx] is None:
                return ""
            val = str(row[idx]).strip()
            return "" if val.lower() == "nan" else val

        missing = [n for n in ("field", "topic", "question", "answer / model answer") if n not in col_idx]
        if missing:
            return Response({"error": f"Missing required column(s): {', '.join(missing)}."}, status=400)

        track_cache, topic_cache, folder_cache = {}, {}, {}
        created_tracks = created_topics = created_folders = created_questions = 0
        skipped = 0

        for row in rows[1:]:
            if not row or all(not str(c or '').strip() for c in row):
                continue

            field_name = col(row, "field")
            topic_title = col(row, "topic")
            question_text = col(row, "question")
            answer_text = col(row, "answer / model answer")
            if not field_name or not topic_title or not question_text or not answer_text:
                skipped += 1
                continue

            if field_name not in track_cache:
                track, created = InterviewTrack.objects.get_or_create(
                    name__iexact=field_name,
                    defaults={"name": field_name, "key": slugify(field_name)[:40] or f"track-{InterviewTrack.objects.count() + 1}"},
                )
                track_cache[field_name] = track
                if created:
                    created_tracks += 1
            track = track_cache[field_name]

            topic_key = (field_name, topic_title)
            if topic_key not in topic_cache:
                topic, created = InterviewTopic.objects.get_or_create(
                    track=track, title=topic_title, defaults={"order": track.topics.count()},
                )
                topic_cache[topic_key] = topic
                if created:
                    created_topics += 1
            topic = topic_cache[topic_key]

            subtopic_title = col(row, "subtopic")
            folder = None
            if subtopic_title:
                folder_key = (topic.id, subtopic_title)
                if folder_key not in folder_cache:
                    folder, created = InterviewFolder.objects.get_or_create(
                        topic=topic, parent=None, title=subtopic_title,
                        defaults={"order": topic.folders.filter(parent__isnull=True).count()},
                    )
                    folder_cache[folder_key] = folder
                    if created:
                        created_folders += 1
                folder = folder_cache[folder_key]

            question_type_raw = col(row, "question type").lower()
            difficulty_raw = col(row, "difficulty").lower()
            question_type = self.QUESTION_TYPE_LABEL_TO_KEY.get(question_type_raw, 'conceptual')
            difficulty = dict(InterviewQuestion.DIFFICULTY_CHOICES).get(
                difficulty_raw.title(), 'Beginner',
            ) if difficulty_raw else 'Beginner'

            InterviewQuestion.objects.create(
                topic=topic, folder=folder,
                external_id=col(row, "question id"),
                question_type=question_type,
                difficulty=difficulty,
                question_text=question_text,
                answer=answer_text,
                follow_up_question=col(row, "follow-up question"),
                follow_up_answer=col(row, "follow-up answer"),
                tools_technologies=col(row, "tools / technologies"),
                key_concepts=col(row, "key concepts / keywords"),
                source_reference=col(row, "source reference"),
                order=(folder.questions.count() if folder else topic.questions.filter(folder__isnull=True).count()),
            )
            created_questions += 1

        return Response({
            "message": "Questions imported",
            "created_tracks": created_tracks,
            "created_topics": created_topics,
            "created_folders": created_folders,
            "created_questions": created_questions,
            "skipped_rows": skipped,
        }, status=201)

    def _read_excel(self, upload):
        import openpyxl
        wb = openpyxl.load_workbook(upload, data_only=True)
        ws = wb["Questions"] if "Questions" in wb.sheetnames else wb.active
        return [[cell.value for cell in row] for row in ws.iter_rows()]

    def _read_csv(self, upload):
        import csv
        import io
        text = upload.read().decode("utf-8-sig")
        return list(csv.reader(io.StringIO(text)))

class AdminInterviewTopicQuestionBulkUploadView(APIView):
    """System Admin: bulk-add questions straight into one Topic (or one of
    its Folders, via optional ?folder_id=) from an uploaded .xlsx/.xls/.csv.
    Same column set and parsing as AdminInterviewQuestionBulkUploadView, but
    Field/Topic/Subtopic columns are ignored if present — the target is
    already fixed by the URL, so admins can upload straight from inside a
    folder without needing the sheet's Field/Topic/Subtopic values to
    exactly match what's already there."""
    permission_classes = [IsAuthenticated]

    QUESTION_TYPE_LABEL_TO_KEY = {label.lower(): key for key, label in InterviewQuestion.QUESTION_TYPE_CHOICES}

    def post(self, request, topic_id):
        if not request.user.is_superuser:
            return Response({"detail": "Admin access required."}, status=status.HTTP_403_FORBIDDEN)
        topic = get_object_or_404(InterviewTopic, id=topic_id)
        folder = None
        folder_id = request.query_params.get('folder_id')
        if folder_id:
            folder = InterviewFolder.objects.filter(id=folder_id, topic=topic).first()
            if not folder:
                return Response({"error": "Folder not found in this topic."}, status=404)

        upload = request.FILES.get("file")
        if not upload:
            return Response({"error": "No file uploaded."}, status=400)

        filename = upload.name.lower()
        try:
            if filename.endswith((".xlsx", ".xls")):
                rows = self._read_excel(upload)
            elif filename.endswith(".csv"):
                rows = self._read_csv(upload)
            else:
                return Response({"error": "Only .xlsx, .xls, or .csv files are supported."}, status=400)
        except Exception as e:
            return Response({"error": f"Could not read file: {e}"}, status=400)

        if not rows:
            return Response({"error": "File is empty."}, status=400)

        header = [str(h or '').strip().lower() for h in rows[0]]
        col_idx = {h: i for i, h in enumerate(header) if h}

        def col(row, name):
            idx = col_idx.get(name)
            if idx is None or idx >= len(row) or row[idx] is None:
                return ""
            val = str(row[idx]).strip()
            return "" if val.lower() == "nan" else val

        missing = [n for n in ("question", "answer / model answer") if n not in col_idx]
        if missing:
            return Response({"error": f"Missing required column(s): {', '.join(missing)}."}, status=400)

        created = []
        skipped = 0
        next_order = folder.questions.count() if folder else topic.questions.filter(folder__isnull=True).count()

        for row in rows[1:]:
            if not row or all(not str(c or '').strip() for c in row):
                continue

            question_text = col(row, "question")
            answer_text = col(row, "answer / model answer")
            if not question_text or not answer_text:
                skipped += 1
                continue

            question_type_raw = col(row, "question type").lower()
            difficulty_raw = col(row, "difficulty").lower()
            question_type = self.QUESTION_TYPE_LABEL_TO_KEY.get(question_type_raw, 'conceptual')
            difficulty = dict(InterviewQuestion.DIFFICULTY_CHOICES).get(
                difficulty_raw.title(), 'Beginner',
            ) if difficulty_raw else 'Beginner'

            q = InterviewQuestion.objects.create(
                topic=topic, folder=folder,
                external_id=col(row, "question id"),
                question_type=question_type,
                difficulty=difficulty,
                question_text=question_text,
                answer=answer_text,
                follow_up_question=col(row, "follow-up question"),
                follow_up_answer=col(row, "follow-up answer"),
                tools_technologies=col(row, "tools / technologies"),
                key_concepts=col(row, "key concepts / keywords"),
                source_reference=col(row, "source reference"),
                order=next_order,
            )
            next_order += 1
            created.append(q)

        return Response({
            "message": "Questions imported",
            "created_questions": len(created),
            "skipped_rows": skipped,
            "questions": [_serialize_interview_question(q) for q in created],
        }, status=201)

    def _read_excel(self, upload):
        import openpyxl
        wb = openpyxl.load_workbook(upload, data_only=True)
        ws = wb["Questions"] if "Questions" in wb.sheetnames else wb.active
        return [[cell.value for cell in row] for row in ws.iter_rows()]

    def _read_csv(self, upload):
        import csv
        import io
        text = upload.read().decode("utf-8-sig")
        return list(csv.reader(io.StringIO(text)))

