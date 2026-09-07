"""Views extracted from the original monolithic apps/learning/views.py
(module: ja/students). Pure code motion — see apps/learning/views/_imports.py
and _shared.py for why this split can't hit a missing-name error."""
from .._imports import *
from .._shared import *


class JABatchListView(APIView):
    """JA: list all batches in their department with student counts."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        profile, err = _ja_guard(request)
        if err:
            return err

        batches = (
            StudentProfile.objects
            .filter(institution=profile.institution, department=profile.department)
            .exclude(batch='')
            .values('batch')
            .annotate(student_count=Count('id'))
            .order_by('-batch')
        )
        return Response({"batches": list(batches)})

    def post(self, request):
        """Create a new (empty) batch — just validates the name is unique in dept."""
        profile, err = _ja_guard(request)
        if err:
            return err

        batch_name = (request.data.get('batch') or '').strip()
        if not batch_name:
            return Response({"detail": "Batch name is required."}, status=status.HTTP_400_BAD_REQUEST)

        # Check if batch already exists in this department
        exists = StudentProfile.objects.filter(
            institution=profile.institution,
            department=profile.department,
            batch=batch_name
        ).exists()
        if exists:
            return Response(
                {"detail": f"Batch '{batch_name}' already exists in this department."},
                status=status.HTTP_400_BAD_REQUEST
            )

        return Response({
            "detail": f"Batch '{batch_name}' is ready. Add students to populate it.",
            "batch": batch_name,
        }, status=status.HTTP_201_CREATED)

class JABatchDetailView(APIView):
    """JA: get students in a batch, or delete the entire batch."""
    permission_classes = [IsAuthenticated]

    def get(self, request, batch_code):
        profile, err = _ja_guard(request)
        if err:
            return err

        students = (
            StudentProfile.objects
            .filter(
                institution=profile.institution,
                department=profile.department,
                batch=batch_code
            )
            .select_related('account', 'mentor')
            .order_by('register_number')
        )

        data = []
        for s in students:
            data.append({
                "id": s.id,
                "register_number": s.register_number,
                "name": s.name,
                "batch": s.batch,
                "section": s.section,
                "personal_email": s.personal_email,
                "mobile_number": s.mobile_number,
                "gender": s.gender,
                "is_active": s.account.is_active if s.account else True,
                "allow_copy_paste": s.allow_copy_paste,
                "mentor_id": s.mentor_id,
                "mentor_name": s.mentor.name if s.mentor_id else None,
                "mentor_faculty_id": s.mentor.faculty_id if s.mentor_id else None,
            })

        return Response({
            "batch": batch_code,
            "students": data,
            "total": len(data),
        })

    def delete(self, request, batch_code):
        """Delete all students in a batch (removes their profiles, accounts, and batch advisor records)."""
        profile, err = _ja_guard(request)
        if err:
            return err

        students = StudentProfile.objects.filter(
            institution=profile.institution,
            department=profile.department,
            batch=batch_code
        ).select_related('account')

        count = students.count()
        user_ids = list(students.values_list('account_id', flat=True))
        students.delete()
        if user_ids:
            User.objects.filter(id__in=user_ids).delete()

        # Delete any batch advisor assignments for this batch
        BatchAdvisor.objects.filter(
            institution=profile.institution,
            department=profile.department,
            batch=batch_code
        ).delete()

        logger.info("JA %s deleted batch '%s' (%d students)", profile.faculty_id, batch_code, count)

        return Response({
            "detail": f"Batch '{batch_code}' deleted successfully ({count} student(s) removed).",
            "deleted_count": count,
        })

class JAStudentCreateView(APIView):
    """JA: add a single student to a batch."""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        profile, err = _ja_guard(request)
        if err:
            return err

        data = request.data
        register_number = (data.get('register_number') or '').strip()
        name = (data.get('name') or '').strip()
        batch = (data.get('batch') or '').strip()

        if not register_number or not name:
            return Response(
                {"detail": "register_number and name are required."},
                status=status.HTTP_400_BAD_REQUEST
            )

        if StudentProfile.objects.filter(register_number=register_number).exists():
            return Response(
                {"detail": f"Student with register number '{register_number}' already exists."},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Create Django User account
        from django.contrib.auth.hashers import make_password

        section = (data.get('section') or '').strip().upper()

        try:
            user = User.objects.create(
                username=register_number,
                password=make_password(None),  # unusable password until first login
                is_active=True,
            )
            student = StudentProfile.objects.create(
                account=user,
                institution=profile.institution,
                department=profile.department,
                register_number=register_number,
                name=name,
                title=data.get('title', name),
                batch=batch,
                section=section,
                personal_email=data.get('personal_email', ''),
                mobile_number=data.get('mobile_number', ''),
                gender=data.get('gender', ''),
            )
        except IntegrityError:
            # Duplicate register_number that slipped past the exists() check (race condition)
            return Response(
                {"detail": f"Student with register number '{register_number}' already exists."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        except Exception:
            logger.exception("Failed to create student %s for JA %s", register_number, profile.faculty_id)
            return Response(
                {"detail": "Failed to create student. Please try again."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        logger.info("JA %s created student %s in batch %s section %s", profile.faculty_id, register_number, batch, section or "—")

        return Response({
            "detail": "Student created successfully.",
            "student": {
                "id": student.id,
                "register_number": student.register_number,
                "name": student.name,
                "batch": student.batch,
                "section": student.section,
            },
        }, status=status.HTTP_201_CREATED)

class JAStudentDeleteView(APIView):
    """JA: remove a student from their department."""
    permission_classes = [IsAuthenticated]

    def delete(self, request, register_number):
        profile, err = _ja_guard(request)
        if err:
            return err

        student = StudentProfile.objects.filter(
            register_number=register_number,
            institution=profile.institution,
            department=profile.department,
        ).select_related('account').first()

        if not student:
            return Response(
                {"detail": "Student not found in your department."},
                status=status.HTTP_404_NOT_FOUND
            )

        student_name = student.name
        user = student.account
        student.delete()
        if user:
            user.delete()

        logger.info("JA %s deleted student %s", profile.faculty_id, register_number)

        return Response({
            "detail": f"Student '{student_name}' ({register_number}) has been removed.",
        })

class JAStudentUpdateView(APIView):
    """JA: update a student's name, mobile number, batch, and/or section.
    PATCH /api/ja/students/<register_number>/update/
    Body: { name?: str, mobile_number?: str, batch?: str, section?: str }
    """
    permission_classes = [IsAuthenticated]

    def patch(self, request, register_number):
        profile, err = _ja_guard(request)
        if err:
            return err

        student = StudentProfile.objects.filter(
            register_number=register_number,
            institution=profile.institution,
            department=profile.department,
        ).first()
        if not student:
            return Response({"detail": "Student not found."}, status=status.HTTP_404_NOT_FOUND)

        fields_updated = []
        name = (request.data.get('name') or '').strip()
        mobile = (request.data.get('mobile_number') or '').strip()

        if name:
            student.name = name
            fields_updated.append('name')
        if 'mobile_number' in request.data:
            student.mobile_number = mobile
            fields_updated.append('mobile_number')
        if 'batch' in request.data:
            student.batch = (request.data.get('batch') or '').strip()
            fields_updated.append('batch')
        if 'section' in request.data:
            raw_section = (request.data.get('section') or '').strip().upper()
            student.section = raw_section
            fields_updated.append('section')

        if not fields_updated:
            return Response({"detail": "Nothing to update."}, status=status.HTTP_400_BAD_REQUEST)

        student.save(update_fields=fields_updated)

        return Response({
            "detail": "Student updated successfully.",
            "register_number": student.register_number,
            "name": student.name,
            "mobile_number": student.mobile_number,
            "batch": student.batch,
            "section": student.section,
        })

class JAStudentMoveView(APIView):
    """JA: move a student to a different batch within the same department."""
    permission_classes = [IsAuthenticated]

    def post(self, request, register_number):
        profile, err = _ja_guard(request)
        if err:
            return err

        new_batch = (request.data.get('batch') or '').strip()
        if not new_batch:
            return Response({"detail": "New batch name is required."}, status=status.HTTP_400_BAD_REQUEST)

        student = StudentProfile.objects.filter(
            register_number=register_number,
            institution=profile.institution,
            department=profile.department,
        ).first()

        if not student:
            return Response(
                {"detail": "Student not found in your department."},
                status=status.HTTP_404_NOT_FOUND
            )

        old_batch = student.batch
        student.batch = new_batch
        student.save(update_fields=['batch'])

        return Response({
            "detail": f"Student moved from batch '{old_batch}' to '{new_batch}'.",
            "register_number": register_number,
            "batch": new_batch,
        })

class JABulkBatchAssignView(APIView):
    """JA: assign multiple students to a batch at once.
    POST /api/ja/students/assign-batch/
    Body: { batch: "23-27", register_numbers: ["REG001", "REG002"] }
    """
    permission_classes = [IsAuthenticated]

    def post(self, request):
        profile, err = _ja_guard(request)
        if err:
            return err

        batch = (request.data.get('batch') or '').strip()
        register_numbers = request.data.get('register_numbers', [])

        if not batch:
            return Response({"detail": "batch is required."}, status=status.HTTP_400_BAD_REQUEST)
        if not register_numbers or not isinstance(register_numbers, list):
            return Response({"detail": "register_numbers must be a non-empty list."}, status=status.HTTP_400_BAD_REQUEST)

        updated = StudentProfile.objects.filter(
            register_number__in=register_numbers,
            institution=profile.institution,
            department=profile.department,
        ).update(batch=batch)

        return Response({
            "detail": f"{updated} student(s) assigned to batch '{batch}'.",
            "updated": updated,
            "batch": batch,
        })

class JABulkImportView(APIView):
    """JA: bulk import students from an uploaded Excel (.xlsx) file.

    Optional POST param: default_batch — used when a row's batch column is empty.
    """
    permission_classes = [IsAuthenticated]

    def post(self, request):
        profile, err = _ja_guard(request)
        if err:
            return err

        excel_file = request.FILES.get('file')
        if not excel_file:
            return Response({"detail": "No file uploaded."}, status=status.HTTP_400_BAD_REQUEST)

        if not excel_file.name.endswith(('.xlsx', '.xls')):
            return Response({"detail": "Only .xlsx or .xls files are supported."}, status=status.HTTP_400_BAD_REQUEST)

        # Optional default batch — used when row's batch column is blank
        default_batch = (request.data.get('default_batch') or '').strip()

        try:
            import openpyxl
        except ImportError:
            return Response(
                {"detail": "openpyxl is not installed on the server."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

        try:
            wb = openpyxl.load_workbook(excel_file, read_only=True, data_only=True)
            ws = wb.active
        except Exception as e:
            return Response({"detail": f"Could not read Excel file: {e}"}, status=status.HTTP_400_BAD_REQUEST)

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return Response({"detail": "Excel file is empty."}, status=status.HTTP_400_BAD_REQUEST)

        # Parse header — batch is now optional if default_batch is provided
        header = [str(h).strip().lower().replace(' ', '_') if h else '' for h in rows[0]]
        required_cols = {'register_number', 'name'}
        if not default_batch:
            required_cols.add('batch')
        missing = required_cols - set(header)
        if missing:
            return Response(
                {"detail": f"Missing required columns: {', '.join(missing)}. See the template."},
                status=status.HTTP_400_BAD_REQUEST
            )

        col_idx = {col: header.index(col) for col in header if col}

        created = []
        created_details = []  # full details for report
        skipped = []
        errors = []

        from django.contrib.auth.hashers import make_password as _make_password

        for row_num, row in enumerate(rows[1:], start=2):
            def get_col(col_name):
                idx = col_idx.get(col_name)
                if idx is None:
                    return ''
                val = row[idx]
                return str(val).strip() if val is not None else ''

            register_number = get_col('register_number')
            name = get_col('name')
            batch = get_col('batch') or default_batch
            section = get_col('section').upper()

            if not register_number or not name:
                errors.append({"row": row_num, "reason": "Missing register_number or name."})
                continue

            if not batch:
                errors.append({"row": row_num, "register_number": register_number, "reason": "No batch specified and no default batch set."})
                continue

            if StudentProfile.objects.filter(register_number=register_number).exists():
                skipped.append({"row": row_num, "register_number": register_number, "reason": "Already exists."})
                continue

            try:
                user = User.objects.create(
                    username=register_number,
                    password=_make_password(None),
                    is_active=True,
                )
                student = StudentProfile.objects.create(
                    account=user,
                    institution=profile.institution,
                    department=profile.department,
                    register_number=register_number,
                    name=name,
                    title=get_col('title') or name,
                    batch=batch,
                    section=section,
                    personal_email=get_col('personal_email'),
                    mobile_number=get_col('mobile_number'),
                    gender=get_col('gender'),
                )
                created.append(register_number)
                created_details.append({
                    "register_number": register_number,
                    "name": name,
                    "batch": batch,
                    "section": section,
                    "title": student.title,
                    "personal_email": student.personal_email,
                    "mobile_number": student.mobile_number,
                    "gender": student.gender,
                    "department": profile.department.name if profile.department else '',
                })
            except Exception as e:
                errors.append({"row": row_num, "register_number": register_number, "reason": str(e)})

        logger.info(
            "JA %s bulk import: %d created, %d skipped, %d errors",
            profile.faculty_id, len(created), len(skipped), len(errors)
        )

        return Response({
            "detail": f"Import complete. {len(created)} created, {len(skipped)} skipped, {len(errors)} errors.",
            "created_count": len(created),
            "skipped_count": len(skipped),
            "error_count": len(errors),
            "created": created,
            "created_details": created_details,
            "skipped": skipped,
            "errors": errors,
        }, status=status.HTTP_200_OK)

class JAExcelTemplateView(APIView):
    """JA: download the Excel template for bulk student import."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        profile, err = _ja_guard(request)
        if err:
            return err

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            return Response(
                {"detail": "openpyxl is not installed on the server."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Students"

        # ── Header row ────────────────────────────────────────────────────────
        headers = [
            "register_number",
            "name",
            "batch",
            "section",
            "title",
            "personal_email",
            "mobile_number",
            "gender",
        ]

        header_fill = PatternFill(start_color="2D6A4F", end_color="2D6A4F", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # ── Sample rows ───────────────────────────────────────────────────────
        sample_rows = [
            ["953623243001", "Arun Kumar", "23-27", "A", "Mr. Arun Kumar", "arun@example.com", "9876543210", "Male"],
            ["953623243002", "Priya Sharma", "23-27", "B", "Ms. Priya Sharma", "priya@example.com", "9876543211", "Female"],
            ["953623243003", "Ravi Raj", "23-27", "A", "Mr. Ravi Raj", "", "", "Male"],
        ]

        sample_fill = PatternFill(start_color="F0FFF4", end_color="F0FFF4", fill_type="solid")
        for row_num, row_data in enumerate(sample_rows, 2):
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.fill = sample_fill
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        # ── Column widths ─────────────────────────────────────────────────────
        col_widths = [20, 25, 12, 10, 30, 30, 16, 10]
        for col_num, width in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = width

        ws.row_dimensions[1].height = 22

        # ── Instructions sheet ────────────────────────────────────────────────
        ws2 = wb.create_sheet("Instructions")
        instructions = [
            ("Column", "Required", "Description"),
            ("register_number", "YES", "Unique student register number (e.g. 953623243001)"),
            ("name", "YES", "Full name of the student"),
            ("batch", "YES", "Batch code (e.g. 23-27)"),
            ("title", "No", "Display title (e.g. Mr. John Doe). Defaults to name if blank."),
            ("section", "No", "Section within the batch: A or B"),
            ("personal_email", "No", "Personal email address"),
            ("mobile_number", "No", "10-digit mobile number"),
            ("gender", "No", "Male / Female / Other"),
        ]
        hdr_font = Font(bold=True)
        for r_idx, row_data in enumerate(instructions, 1):
            for c_idx, val in enumerate(row_data, 1):
                cell = ws2.cell(row=r_idx, column=c_idx, value=val)
                if r_idx == 1:
                    cell.font = hdr_font
                    cell.fill = PatternFill(start_color="D8F3DC", end_color="D8F3DC", fill_type="solid")
                cell.border = thin_border

        ws2.column_dimensions['A'].width = 22
        ws2.column_dimensions['B'].width = 12
        ws2.column_dimensions['C'].width = 55

        # ── Stream response ───────────────────────────────────────────────────
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        dept_code = profile.department.code if profile.department else "dept"
        filename = f"student_import_template_{dept_code}.xlsx"

        response = HttpResponse(
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

class JAStudentListView(APIView):
    """JA: list all students in their department with optional batch/search filter."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        profile, err = _ja_guard(request)
        if err:
            return err

        students = StudentProfile.objects.filter(
            institution=profile.institution,
            department=profile.department,
        ).select_related('account', 'mentor').order_by('batch', 'register_number')

        batch = request.query_params.get('batch')
        section_param = (request.query_params.get('section') or '').strip()
        search = request.query_params.get('search', '').strip()

        if batch:
            students = students.filter(batch=batch)
        if section_param == '__none__':
            students = students.filter(section='')
        elif section_param:
            students = students.filter(section=section_param.upper())
        if search:
            students = students.filter(
                Q(name__icontains=search) | Q(register_number__icontains=search)
            )

        data = []
        for s in students:
            data.append({
                "id": s.id,
                "register_number": s.register_number,
                "name": s.name,
                "batch": s.batch,
                "section": s.section,
                "personal_email": s.personal_email,
                "mobile_number": s.mobile_number,
                "gender": s.gender,
                "is_active": s.account.is_active if s.account else True,
                "mentor_id": s.mentor_id,
                "mentor_name": s.mentor.name if s.mentor_id else None,
                "mentor_faculty_id": s.mentor.faculty_id if s.mentor_id else None,
            })

        return Response({"students": data, "total": len(data)})

class JAImportReportView(APIView):
    """JA: generate and download an Excel report of a list of students (by register numbers)."""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        profile, err = _ja_guard(request)
        if err:
            return err

        register_numbers = request.data.get('register_numbers', [])
        report_title = request.data.get('title', 'Student Report')

        if not register_numbers:
            return Response({"detail": "No register numbers provided."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            return Response({"detail": "openpyxl not installed."}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        students = StudentProfile.objects.filter(
            register_number__in=register_numbers,
            institution=profile.institution,
            department=profile.department,
        ).order_by('batch', 'register_number')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Students"

        # ── Styles ────────────────────────────────────────────────────────────
        hdr_fill  = PatternFill(start_color="2D6A4F", end_color="2D6A4F", fill_type="solid")
        hdr_font  = Font(bold=True, color="FFFFFF", size=11)
        meta_font = Font(bold=True, size=12, color="1A3C2A")
        center    = Alignment(horizontal="center", vertical="center")
        thin      = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        row_fill  = PatternFill(start_color="F0FFF4", end_color="F0FFF4", fill_type="solid")

        # ── Title rows ────────────────────────────────────────────────────────
        dept_name = profile.department.name if profile.department else ''
        inst_name = profile.institution.name if profile.institution else ''

        ws.merge_cells('A1:H1')
        ws['A1'] = inst_name
        ws['A1'].font = Font(bold=True, size=14, color="1A3C2A")
        ws['A1'].alignment = center

        ws.merge_cells('A2:H2')
        ws['A2'] = f"{dept_name} — {report_title}"
        ws['A2'].font = meta_font
        ws['A2'].alignment = center

        ws.merge_cells('A3:H3')
        from django.utils import timezone as tz
        ws['A3'] = f"Generated: {tz.now().strftime('%d %b %Y, %I:%M %p')}"
        ws['A3'].font = Font(size=10, color="6B7280")
        ws['A3'].alignment = center

        ws.row_dimensions[1].height = 22
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 16

        # ── Header row ────────────────────────────────────────────────────────
        headers = ['#', 'Register Number', 'Name', 'Batch', 'Title', 'Email', 'Mobile', 'Gender']
        for col_num, h in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_num, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = center
            cell.border = thin
        ws.row_dimensions[5].height = 20

        # ── Data rows ─────────────────────────────────────────────────────────
        for idx, student in enumerate(students, 1):
            row_num = idx + 5
            values = [
                idx,
                student.register_number,
                student.name,
                student.batch,
                student.title,
                student.personal_email,
                student.mobile_number,
                student.gender,
            ]
            for col_num, val in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.border = thin
                cell.alignment = Alignment(vertical="center")
                if idx % 2 == 0:
                    cell.fill = row_fill

        # ── Column widths ─────────────────────────────────────────────────────
        col_widths = [5, 20, 28, 10, 32, 30, 16, 10]
        for col_num, width in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = width

        # ── Summary row ───────────────────────────────────────────────────────
        summary_row = students.count() + 7
        ws.cell(row=summary_row, column=1, value=f"Total: {students.count()} student(s)").font = Font(bold=True, color="2D6A4F")

        # ── Stream ────────────────────────────────────────────────────────────
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        dept_code = profile.department.code if profile.department else 'dept'
        from django.utils import timezone as tz2
        ts = tz2.now().strftime('%Y%m%d_%H%M')
        filename = f"students_{dept_code}_{ts}.xlsx"

        response = HttpResponse(
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

class JABulkSectionAssignView(APIView):
    """JA: assign multiple students to a section at once.
    POST /api/ja/students/assign-section/
    Body: { section: "A", register_numbers: ["REG001", "REG002"] }
    """
    permission_classes = [IsAuthenticated]

    def post(self, request):
        profile, err = _ja_guard(request)
        if err:
            return err

        section = (request.data.get('section') or '').strip().upper()
        register_numbers = request.data.get('register_numbers', [])

        if not section:
            return Response({"detail": "section is required."}, status=status.HTTP_400_BAD_REQUEST)
        if not register_numbers or not isinstance(register_numbers, list):
            return Response({"detail": "register_numbers must be a non-empty list."}, status=status.HTTP_400_BAD_REQUEST)

        updated = StudentProfile.objects.filter(
            register_number__in=register_numbers,
            institution=profile.institution,
            department=profile.department,
        ).update(section=section)

        return Response({
            "detail": f"{updated} student(s) assigned to section '{section}'.",
            "updated": updated,
            "section": section,
        })

